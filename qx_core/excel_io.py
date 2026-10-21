from __future__ import annotations

import re
import uuid
from collections.abc import Mapping
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .models import DormRoom, Student


def load_students_from_excel(path: Path) -> list[Student]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    students: list[Student] = []

    for row in sheet.iter_rows(min_row=1, values_only=True):
        if not row or all(cell is None for cell in row[:5]):
            continue

        raw_index, raw_name, raw_hometown, raw_class, raw_area = (list(row) + [None] * 5)[:5]
        if str(raw_index).strip() in {"序号", "编号"}:
            continue

        name = "" if raw_name is None else str(raw_name).strip()
        if not name:
            continue

        try:
            index = int(raw_index)
        except (TypeError, ValueError):
            index = len(students) + 1

        hometown = "" if raw_hometown is None else str(raw_hometown).strip()
        class_name = "" if raw_class is None else str(raw_class).strip()
        students.append(Student(index=index, name=name, hometown=hometown, original_index=index, class_name=class_name, source_area=str(raw_area or "").strip()))

    return students


def normalized_match_text(value: str) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip().lower()


def excel_cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_dorm_rooms_from_excel(path: Path, students: list[Student]) -> tuple[list[DormRoom], int]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    students_by_exact: dict[tuple[str, str], list[Student]] = {}
    students_by_name: dict[str, list[Student]] = {}
    for student in students:
        name_key = normalized_match_text(student.name)
        hometown_key = normalized_match_text(student.hometown)
        if not name_key:
            continue
        students_by_name.setdefault(name_key, []).append(student)
        students_by_exact.setdefault((name_key, hometown_key), []).append(student)

    rooms: list[DormRoom] = []
    room_by_number: dict[str, DormRoom] = {}
    current_room_number = ""
    matched_count = 0
    used_keys: set[str] = set()

    for row in sheet.iter_rows(min_row=1, values_only=True):
        values = [excel_cell_text(cell) for cell in row[:4]]
        values.extend([""] * (4 - len(values)))
        if not any(values):
            continue
        if any(value in {"宿舍号", "床号", "姓名", "籍贯"} for value in values):
            continue
        room_number, bed_text, student_name, hometown = values
        if room_number and not any((bed_text, student_name, hometown)):
            continue
        if not room_number:
            room_number = current_room_number
        if not room_number:
            continue
        current_room_number = room_number

        room = room_by_number.get(room_number)
        if room is None:
            room = DormRoom(room_number, [])
            rooms.append(room)
            room_by_number[room_number] = room

        try:
            bed_number = int(re.sub(r"\D+", "", bed_text)) if bed_text else len(room.beds) + 1
        except ValueError:
            bed_number = len(room.beds) + 1
        bed_index = max(0, bed_number - 1)
        while len(room.beds) <= bed_index:
            room.beds.append(None)

        student_key = ""
        name_key = normalized_match_text(student_name)
        hometown_key = normalized_match_text(hometown)
        candidates = students_by_exact.get((name_key, hometown_key), [])
        if not candidates and name_key:
            candidates = students_by_name.get(name_key, [])
        for candidate in candidates:
            if not candidate.dorm_key:
                candidate.dorm_key = str(uuid.uuid4())
            if candidate.dorm_key not in used_keys:
                student_key = candidate.dorm_key
                used_keys.add(student_key)
                matched_count += 1
                break
        room.beds[bed_index] = student_key or None

    return rooms, matched_count


def export_students_to_excel(path: Path, students: list[Student]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "学生信息"
    sheet.append(["序号", "姓名", "籍贯", "班级", "送生区"])
    for student in students:
        sheet.append([student.index, student.name, student.hometown, student.class_name, student.source_area])
    workbook.save(path)


def export_dorm_roster_to_excel(
    path: Path,
    rooms: list[DormRoom],
    students_by_key: dict[str, Student],
    title: str = "宿舍安排表",
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "宿舍安排表"
    sheet.merge_cells("A1:E1")
    sheet["A1"] = title
    sheet["A1"].font = Font(name="宋体", size=22, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 38

    headers = ["宿舍号", "床号", "姓名", "籍贯", "送生区"]
    sheet.append(headers)
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="F2F2F2")
    for cell in sheet[2]:
        cell.font = Font(name="宋体", size=14, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        cell.fill = header_fill
    sheet.row_dimensions[2].height = 30

    current_row = 3
    for room in rooms:
        if not room.beds:
            continue
        start_row = current_row
        for bed_index, student_key in enumerate(room.beds, start=1):
            student = students_by_key.get(student_key or "")
            sheet.cell(current_row, 1, room.number)
            sheet.cell(current_row, 2, f"{bed_index:02d}")
            sheet.cell(current_row, 3, student.name if student else "")
            sheet.cell(current_row, 4, student.hometown if student else "")
            sheet.cell(current_row, 5, student.source_area if student else "")
            for column in range(1, 6):
                cell = sheet.cell(current_row, column)
                cell.font = Font(name="宋体", size=13)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border
            sheet.row_dimensions[current_row].height = 26
            current_row += 1
        end_row = current_row - 1
        if end_row > start_row:
            sheet.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
        room_cell = sheet.cell(start_row, 1)
        room_cell.alignment = Alignment(horizontal="center", vertical="center")
        room_cell.font = Font(name="宋体", size=14)
        room_cell.border = border

    sheet.column_dimensions["A"].width = 10
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["C"].width = 22
    sheet.column_dimensions["D"].width = 26
    sheet.column_dimensions["E"].width = 22
    for row in sheet.iter_rows(min_row=1, max_row=max(2, current_row - 1), min_col=1, max_col=5):
        for cell in row:
            cell.border = border
    workbook.save(path)


PHOTO_COVERAGE_HEADERS = ["序号", "姓名", "籍贯", "班级", "照片状态", "匹配方式", "送生区"]
PHOTO_MATCH_STATUS_LABELS = {
    "strict": "精确",
    "fuzzy": "模糊",
    "manual": "手动",
    "missing": "缺照片",
    "unmatched": "未匹配",
    "ambiguous": "多重匹配",
    "duplicate": "重复",
    "rejected": "已拒绝",
}


def photo_match_status_label(value: object) -> str:
    text = str(value or "")
    return PHOTO_MATCH_STATUS_LABELS.get(text, text)


def photo_coverage_row(student: Mapping[str, object], translate_match_status: bool = False) -> list[object]:
    has_photo = bool(student.get("photo_path"))
    match_status = student.get("photo_match_status") or ""
    if translate_match_status:
        match_status = photo_match_status_label(match_status)
    return [
        student.get("index") or "",
        student.get("name") or "",
        student.get("hometown") or "",
        student.get("class_name") or "",
        "已绑定" if has_photo else "缺照片",
        match_status,
        student.get("source_area") or "",
    ]


def style_photo_table(
    sheet,
    header_row: int,
    header_fill: PatternFill,
    border: Border,
    widths: list[int],
) -> None:
    max_column = len(widths)
    for cell in sheet[header_row]:
        cell.font = Font(name="宋体", size=12, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = header_fill
        cell.border = border
    for row in sheet.iter_rows(min_row=header_row + 1, max_row=sheet.max_row, min_col=1, max_col=max_column):
        for cell in row:
            cell.font = Font(name="宋体", size=11)
            cell.alignment = Alignment(vertical="center")
            cell.border = border
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + column)].width = width
    sheet.freeze_panes = f"A{header_row + 1}"


def write_photo_coverage_table(
    sheet,
    students: list[Mapping[str, object]],
    header_row: int,
    header_fill: PatternFill,
    border: Border,
    translate_match_status: bool = False,
) -> None:
    for column, header in enumerate(PHOTO_COVERAGE_HEADERS, start=1):
        sheet.cell(header_row, column, header)
    for row_index, student in enumerate(students, start=header_row + 1):
        for column, value in enumerate(photo_coverage_row(student, translate_match_status), start=1):
            sheet.cell(row_index, column, value)
    style_photo_table(sheet, header_row, header_fill, border, [10, 16, 20, 18, 12, 14, 22])


def write_missing_photo_sheet(
    workbook: Workbook,
    students: list[Mapping[str, object]],
    header_fill: PatternFill,
    border: Border,
) -> None:
    missing_sheet = workbook.create_sheet("缺照片学生")
    missing_headers = PHOTO_COVERAGE_HEADERS[:4] + ["送生区"]
    for column, header in enumerate(missing_headers, start=1):
        missing_sheet.cell(1, column, header)
    missing_row = 2
    for student in students:
        if student.get("photo_path"):
            continue
        values = [
            student.get("index") or "",
            student.get("name") or "",
            student.get("hometown") or "",
            student.get("class_name") or "",
            student.get("source_area") or "",
        ]
        for column, value in enumerate(values, start=1):
            missing_sheet.cell(missing_row, column, value)
        missing_row += 1
    style_photo_table(missing_sheet, 1, header_fill, border, [10, 16, 20, 18, 22])


def student_index_value(student: Mapping[str, object]) -> int | None:
    value = student.get("index")
    if isinstance(value, int):
        return value
    try:
        return int(str(value))
    except (TypeError, ValueError):
        return None


def write_photo_sequence_gap_sheet(
    workbook: Workbook,
    missing_sequence_numbers: list[int],
    students: list[Mapping[str, object]],
    header_fill: PatternFill,
    border: Border,
) -> None:
    sheet = workbook.create_sheet("文件序号缺口")
    headers = ["缺失序号", "名单学生", "学生ID", "籍贯", "班级", "照片状态", "匹配方式", "送生区"]
    for column, header in enumerate(headers, start=1):
        sheet.cell(1, column, header)

    students_by_index = {
        index: student
        for student in students
        if (index := student_index_value(student)) is not None
    }
    for row_index, number in enumerate(missing_sequence_numbers, start=2):
        student = students_by_index.get(number)
        has_photo = bool(student and student.get("photo_path"))
        values = [
            f"{number:02d}",
            student.get("name") if student else "",
            student.get("id") if student else "",
            student.get("hometown") if student else "",
            student.get("class_name") if student else "",
            "已绑定" if has_photo else "缺照片" if student else "名单无此序号",
            photo_match_status_label(student.get("photo_match_status")) if student else "",
            student.get("source_area") if student else "",
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value)
    style_photo_table(sheet, 1, header_fill, border, [10, 16, 12, 20, 18, 12, 14, 22])


def write_photo_image_summary_sheet(
    workbook: Workbook,
    results: list[Mapping[str, object]],
    header_fill: PatternFill,
    border: Border,
) -> None:
    sheet = workbook.create_sheet("图片规格汇总")
    headers = ["图片尺寸", "图片格式", "数量", "示例文件"]
    for column, header in enumerate(headers, start=1):
        sheet.cell(1, column, header)

    grouped: dict[tuple[str, str], dict[str, object]] = {}
    for item in results:
        width = item.get("image_width")
        height = item.get("image_height")
        size = f"{width} x {height}" if width and height else "无法读取"
        image_format = str(item.get("image_format") or "未知")
        key = (size, image_format)
        group = grouped.setdefault(key, {"count": 0, "examples": []})
        group["count"] = int(group["count"]) + 1
        examples = group["examples"]
        if isinstance(examples, list) and len(examples) < 3:
            examples.append(str(item.get("filename") or ""))

    rows = sorted(
        grouped.items(),
        key=lambda entry: (-int(entry[1]["count"]), entry[0][0], entry[0][1]),
    )
    for row_index, ((size, image_format), data) in enumerate(rows, start=2):
        examples = data["examples"]
        sheet.cell(row_index, 1, size)
        sheet.cell(row_index, 2, image_format)
        sheet.cell(row_index, 3, data["count"])
        sheet.cell(row_index, 4, "、".join(examples) if isinstance(examples, list) else "")
    style_photo_table(sheet, 1, header_fill, border, [16, 12, 10, 42])


def export_photo_import_report_to_excel(
    path: Path,
    matched_count: int,
    unmatched_count: int,
    sequence_numbers: list[int],
    missing_sequence_numbers: list[int],
    results: list[Mapping[str, object]],
    students: list[Mapping[str, object]] | None = None,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "照片导入核对"

    title = "照片导入核对清单"
    sheet.merge_cells("A1:K1")
    sheet["A1"] = title
    sheet["A1"].font = Font(name="宋体", size=18, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 32

    summary_rows = [
        ("已匹配", matched_count),
        ("未匹配/需处理", unmatched_count),
        ("已有文件序号", "、".join(f"{number:02d}" for number in sequence_numbers)),
        ("缺失文件序号", "、".join(f"{number:02d}" for number in missing_sequence_numbers) or "无"),
    ]
    for row_index, (label, value) in enumerate(summary_rows, start=2):
        sheet.cell(row_index, 1, label)
        sheet.cell(row_index, 2, value)

    headers = [
        "文件序号",
        "文件名",
        "状态",
        "说明",
        "匹配学生",
        "建议序号",
        "建议学生",
        "图片尺寸",
        "图片格式",
        "上传文件ID",
        "下载地址",
    ]
    header_row = 7
    for column, header in enumerate(headers, start=1):
        sheet.cell(header_row, column, header)

    status_labels = {
        "strict": "精确",
        "fuzzy": "模糊",
        "manual": "手动",
        "unmatched": "未匹配",
        "ambiguous": "多重匹配",
        "duplicate": "重复",
        "rejected": "已拒绝",
    }
    for row_index, item in enumerate(results, start=header_row + 1):
        sequence_number = item.get("sequence_number")
        suggested_index = item.get("suggested_student_index")
        image_width = item.get("image_width")
        image_height = item.get("image_height")
        image_size = f"{image_width} x {image_height}" if image_width and image_height else ""
        values = [
            f"{sequence_number:02d}" if isinstance(sequence_number, int) else "",
            item.get("filename") or "",
            status_labels.get(str(item.get("status") or ""), item.get("status") or ""),
            item.get("message") or "",
            item.get("student_name") or "",
            f"{suggested_index:02d}" if isinstance(suggested_index, int) else "",
            item.get("suggested_student_name") or "",
            image_size,
            item.get("image_format") or "",
            item.get("file_id") or "",
            item.get("download_url") or "",
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value)

    thin = Side(style="thin", color="D9E2EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="EAF2FF")
    summary_fill = PatternFill("solid", fgColor="F8FAFC")
    for row in sheet.iter_rows(min_row=2, max_row=5, min_col=1, max_col=2):
        for cell in row:
            cell.font = Font(name="宋体", size=12)
            cell.alignment = Alignment(vertical="center")
            cell.fill = summary_fill
            cell.border = border
    for cell in sheet[header_row]:
        cell.font = Font(name="宋体", size=12, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = header_fill
        cell.border = border
    for row in sheet.iter_rows(min_row=header_row + 1, max_row=max(header_row + 1, header_row + len(results)), min_col=1, max_col=len(headers)):
        for cell in row:
            cell.font = Font(name="宋体", size=11)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border

    widths = [10, 28, 10, 26, 16, 10, 16, 14, 12, 34, 36]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A8"

    students = students or []
    write_photo_image_summary_sheet(workbook, results, header_fill, border)
    write_photo_sequence_gap_sheet(workbook, missing_sequence_numbers, students, header_fill, border)
    coverage_sheet = workbook.create_sheet("学生照片覆盖")
    write_photo_coverage_table(coverage_sheet, students, 1, header_fill, border)
    write_missing_photo_sheet(workbook, students, header_fill, border)
    workbook.save(path)


def export_photo_status_report_to_excel(path: Path, students: list[Mapping[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "照片状态清单"

    sheet.merge_cells("A1:G1")
    sheet["A1"] = "学生照片状态清单"
    sheet["A1"].font = Font(name="宋体", size=18, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 32

    bound_count = sum(1 for student in students if student.get("photo_path"))
    summary_rows = [
        ("总人数", len(students)),
        ("已绑定照片", bound_count),
        ("缺照片", len(students) - bound_count),
    ]
    thin = Side(style="thin", color="D9E2EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="EAF2FF")
    summary_fill = PatternFill("solid", fgColor="F8FAFC")
    for row_index, (label, value) in enumerate(summary_rows, start=2):
        sheet.cell(row_index, 1, label)
        sheet.cell(row_index, 2, value)
    for row in sheet.iter_rows(min_row=2, max_row=4, min_col=1, max_col=2):
        for cell in row:
            cell.font = Font(name="宋体", size=12)
            cell.alignment = Alignment(vertical="center")
            cell.fill = summary_fill
            cell.border = border

    write_photo_coverage_table(sheet, students, 6, header_fill, border, translate_match_status=True)
    write_missing_photo_sheet(workbook, students, header_fill, border)
    workbook.save(path)
