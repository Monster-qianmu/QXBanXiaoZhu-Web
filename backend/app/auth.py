from __future__ import annotations

import hashlib
import hmac
import io
import re
import secrets
import threading
import time
from collections import defaultdict, deque
from datetime import datetime, timedelta

from fastapi import APIRouter, Depends, File, HTTPException, Request, Response, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field as PydanticField
from sqlalchemy.exc import IntegrityError
from sqlmodel import Session, select

from .db import get_session
from .models import LoginSessionRecord, ProjectRecord, RegistrationRequestRecord, UserRecord
from .registration import approve_due_registrations, normalize_phone

router = APIRouter(prefix="/api/v1")
COOKIE_NAME = "qxban_session"
PUBLIC_PATHS = {"/api/v1/health", "/api/v1/auth/status", "/api/v1/auth/setup", "/api/v1/auth/login", "/api/v1/auth/registration-requests"}
_attempts: dict[str, deque] = defaultdict(deque)
_attempt_lock = threading.Lock()


class Credentials(BaseModel):
    username: str = PydanticField(min_length=1, max_length=64)
    password: str = PydanticField(min_length=1, max_length=128)


class CreateUser(Credentials):
    display_name: str = PydanticField(default="", max_length=100)


class ChangePassword(BaseModel):
    current_password: str = PydanticField(max_length=128)
    new_password: str = PydanticField(max_length=128)


class ResetPassword(BaseModel):
    password: str = PydanticField(max_length=128)


class UserPatch(BaseModel):
    is_active: bool


def user_out(user: UserRecord) -> dict:
    return {key: getattr(user, key) for key in ("id", "username", "phone", "display_name", "role", "is_active", "must_change_password", "created_at")}


def normalize_username(value: str) -> str:
    value = value.strip().lower()
    if not re.fullmatch(r"[a-z0-9_.@-]{3,64}", value):
        raise HTTPException(400, "账号须为 3–64 位字母、数字或 _ . @ -")
    return value


def validate_password(value: str) -> None:
    if len(value) < 8 or len(value) > 128 or not value.strip():
        raise HTTPException(400, "密码须为 8–128 个字符")


def hash_password(value: str) -> str:
    salt = secrets.token_hex(16)
    digest = hashlib.scrypt(value.encode(), salt=salt.encode(), n=16384, r=8, p=1).hex()
    return f"scrypt${salt}${digest}"


def verify_password(value: str, encoded: str) -> bool:
    try:
        algorithm, salt, expected = encoded.split("$")
        if algorithm != "scrypt":
            return False
        actual = hashlib.scrypt(value.encode(), salt=salt.encode(), n=16384, r=8, p=1).hex()
        return hmac.compare_digest(actual, expected)
    except (ValueError, TypeError):
        return False


def authorize_request(request: Request, session: Session = Depends(get_session)) -> None:
    if request.method in {"POST", "PATCH", "PUT", "DELETE"}:
        origin = request.headers.get("origin")
        if origin and origin != str(request.base_url).rstrip("/"):
            raise HTTPException(403, "请求来源不受信任，请从系统页面操作")
    if request.url.path in PUBLIC_PATHS:
        return
    token = request.cookies.get(COOKIE_NAME, "")
    record = session.get(LoginSessionRecord, hashlib.sha256(token.encode()).hexdigest()) if token else None
    if record is None or record.expires_at <= datetime.utcnow():
        raise HTTPException(401, "请先登录")
    user = session.get(UserRecord, record.user_id)
    if user is None or not user.is_active:
        raise HTTPException(401, "账号已停用，请联系管理员")
    session.info["user"] = user
    if user.must_change_password and request.url.path not in {"/api/v1/auth/me", "/api/v1/auth/logout", "/api/v1/auth/change-password"}:
        raise HTTPException(403, "请先修改初始密码")


def current_user(session: Session = Depends(get_session)) -> UserRecord:
    user = session.info.get("user")
    if user is None:
        raise HTTPException(401, "请先登录")
    return user


def admin_user(user: UserRecord = Depends(current_user)) -> UserRecord:
    if user.role != "admin":
        raise HTTPException(403, "仅管理员可以管理账号")
    return user


def start_session(user: UserRecord, session: Session, response: Response, request: Request) -> None:
    token = secrets.token_urlsafe(40)
    session.add(LoginSessionRecord(id=hashlib.sha256(token.encode()).hexdigest(), user_id=user.id, expires_at=datetime.utcnow() + timedelta(hours=12)))
    session.commit()
    response.set_cookie(COOKIE_NAME, token, max_age=43200, httponly=True, samesite="strict", secure=request.url.scheme == "https", path="/")


def revoke_sessions(session: Session, user_id: int) -> None:
    for record in session.exec(select(LoginSessionRecord).where(LoginSessionRecord.user_id == user_id)):
        session.delete(record)


@router.get("/auth/status")
def auth_status(session: Session = Depends(get_session)):
    return {"needs_setup": session.exec(select(UserRecord.id).where(UserRecord.role == "admin")).first() is None}


@router.post("/auth/setup")
def setup(payload: CreateUser, response: Response, request: Request, session: Session = Depends(get_session)):
    username = normalize_username(payload.username)
    validate_password(payload.password)
    password_hash = hash_password(payload.password)
    # Serialise first-admin creation across concurrent requests/processes.
    session.connection().exec_driver_sql("BEGIN IMMEDIATE")
    if session.exec(select(UserRecord.id).where(UserRecord.role == "admin")).first() is not None:
        session.rollback()
        raise HTTPException(409, "管理员已初始化，请登录")
    user = UserRecord(username=username, display_name=payload.display_name.strip() or username, password_hash=password_hash, role="admin")
    session.add(user)
    session.flush()
    for project in session.exec(select(ProjectRecord).where(ProjectRecord.owner_id == None)):
        project.owner_id = user.id
        session.add(project)
    session.commit()
    session.refresh(user)
    start_session(user, session, response, request)
    return user_out(user)


@router.post("/auth/login")
def login(payload: Credentials, response: Response, request: Request, session: Session = Depends(get_session)):
    login_name = payload.username.strip().lower()
    try:
        login_name = normalize_phone(login_name)
    except ValueError:
        pass
    key = (request.client.host if request.client else "unknown") + ":" + login_name
    with _attempt_lock:
        now = time.monotonic()
        attempts = _attempts[key]
        while attempts and attempts[0] < now - 300:
            attempts.popleft()
        if len(attempts) >= 10:
            raise HTTPException(429, "尝试次数过多，请 5 分钟后再试")
        attempts.append(now)
    approve_due_registrations()
    user = session.exec(select(UserRecord).where(UserRecord.username == login_name)).first()
    if user is None:
        application = session.exec(select(RegistrationRequestRecord).where(
            RegistrationRequestRecord.phone == login_name,
            RegistrationRequestRecord.status.in_(["pending", "rejected"]),
        ).order_by(RegistrationRequestRecord.id.desc())).first()
        if application is not None and verify_password(payload.password, application.password_hash):
            if application.status == "pending":
                raise HTTPException(403, "账号申请正在等待管理员审批，提交满 24 小时未处理将自动同意")
            raise HTTPException(403, f"账号申请未通过：{application.review_note}，可修改信息后重新申请")
    if user is None or not verify_password(payload.password, user.password_hash) or not user.is_active:
        raise HTTPException(401, "账号或密码错误，或账号已停用")
    with _attempt_lock:
        _attempts.pop(key, None)
    start_session(user, session, response, request)
    return user_out(user)


@router.get("/auth/me")
def me(user: UserRecord = Depends(current_user)):
    return user_out(user)


@router.post("/auth/logout")
def logout(request: Request, response: Response, session: Session = Depends(get_session)):
    token = request.cookies.get(COOKIE_NAME, "")
    record = session.get(LoginSessionRecord, hashlib.sha256(token.encode()).hexdigest())
    if record:
        session.delete(record)
        session.commit()
    response.delete_cookie(COOKIE_NAME, path="/")
    return {"ok": True}


@router.post("/auth/change-password")
def change_password(payload: ChangePassword, response: Response, request: Request, user: UserRecord = Depends(current_user), session: Session = Depends(get_session)):
    if not verify_password(payload.current_password, user.password_hash):
        raise HTTPException(400, "当前密码错误")
    validate_password(payload.new_password)
    if payload.new_password == payload.current_password:
        raise HTTPException(400, "新密码不能与当前密码相同")
    user.password_hash = hash_password(payload.new_password)
    user.must_change_password = False
    session.add(user)
    revoke_sessions(session, user.id)
    session.commit()
    start_session(user, session, response, request)
    return user_out(user)


@router.get("/admin/users")
def list_users(_admin: UserRecord = Depends(admin_user), session: Session = Depends(get_session)):
    return [user_out(user) for user in session.exec(select(UserRecord).order_by(UserRecord.id))]


@router.post("/admin/users")
def create_user(payload: CreateUser, _admin: UserRecord = Depends(admin_user), session: Session = Depends(get_session)):
    username = normalize_username(payload.username)
    validate_password(payload.password)
    user = UserRecord(username=username, display_name=payload.display_name.strip() or username, password_hash=hash_password(payload.password), must_change_password=True)
    session.add(user)
    try:
        session.commit()
    except IntegrityError:
        session.rollback()
        raise HTTPException(409, "账号已存在")
    session.refresh(user)
    return user_out(user)


@router.get("/admin/users/template")
def user_template(_admin: UserRecord = Depends(admin_user)):
    workbook = Workbook()
    workbook.active.title = "账号导入"
    workbook.active.append(["账号", "姓名", "初始密码"])
    workbook.active.column_dimensions["A"].width = 24
    workbook.active.column_dimensions["B"].width = 20
    workbook.active.column_dimensions["C"].width = 24
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=users-template.xlsx"})


@router.post("/admin/users/import")
def import_users(file: UploadFile = File(...), _admin: UserRecord = Depends(admin_user), session: Session = Depends(get_session)):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(400, "请选择 .xlsx 账号表")
    data = file.file.read(2 * 1024 * 1024 + 1)
    if len(data) > 2 * 1024 * 1024:
        raise HTTPException(400, "账号表不能超过 2 MB")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
    except Exception:
        raise HTTPException(400, "无法读取账号表，请检查文件格式")
    if not rows or list(rows[0][:3]) != ["账号", "姓名", "初始密码"]:
        raise HTTPException(400, "表头须为：账号、姓名、初始密码，请下载账号样表")
    if len(rows) > 501:
        raise HTTPException(400, "一次最多导入 500 个账号")
    existing = set(session.exec(select(UserRecord.username)))
    pending = []
    errors = []
    for index, row in enumerate(rows[1:], 2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        raw_username, raw_name, raw_password = (list(row) + [None] * 3)[:3]
        try:
            username = normalize_username(str(raw_username or ""))
            password = str(raw_password or "")
            validate_password(password)
            if username in existing:
                raise HTTPException(400, "账号已存在或在表内重复")
            if len(str(raw_name or "")) > 100:
                raise HTTPException(400, "姓名不能超过 100 个字符")
            existing.add(username)
            pending.append(UserRecord(username=username, display_name=str(raw_name or username).strip(), password_hash=hash_password(password), must_change_password=True))
        except HTTPException as exc:
            errors.append({"row": index, "message": exc.detail})
    if errors:
        raise HTTPException(400, {"message": "导入未写入任何账号，请修正以下行", "errors": errors})
    if not pending:
        raise HTTPException(400, "账号表中没有可导入的数据")
    session.add_all(pending)
    try:
        session.commit()
    except IntegrityError:
        session.rollback()
        raise HTTPException(409, "账号已存在，请刷新后重新导入")
    return {"imported_count": len(pending)}


def ordinary_user(session: Session, user_id: int) -> UserRecord:
    user = session.get(UserRecord, user_id)
    if user is None:
        raise HTTPException(404, "账号不存在")
    if user.role != "user":
        raise HTTPException(403, "此操作仅适用于普通用户")
    return user


@router.post("/admin/users/{user_id}/reset-password")
def reset_password(user_id: int, payload: ResetPassword, _admin: UserRecord = Depends(admin_user), session: Session = Depends(get_session)):
    user = ordinary_user(session, user_id)
    validate_password(payload.password)
    user.password_hash = hash_password(payload.password)
    user.must_change_password = True
    session.add(user)
    revoke_sessions(session, user.id)
    session.commit()
    return {"ok": True}


@router.patch("/admin/users/{user_id}")
def patch_user(user_id: int, payload: UserPatch, _admin: UserRecord = Depends(admin_user), session: Session = Depends(get_session)):
    user = ordinary_user(session, user_id)
    user.is_active = payload.is_active
    session.add(user)
    if not user.is_active:
        revoke_sessions(session, user.id)
    session.commit()
    return user_out(user)
