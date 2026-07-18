from datetime import timedelta
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from PIL import Image


def write_students_excel(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["序号", "姓名", "籍贯", "班级"])
    sheet.append([1, "张三", "广州", "一班"])
    sheet.append([2, "李四", "深圳", "一班"])
    workbook.save(path)


def write_numbered_students_excel(path: Path, names: list[str]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["序号", "姓名", "籍贯", "班级"])
    for index, name in enumerate(names, start=1):
        sheet.append([index, name, "", "一班"])
    workbook.save(path)


def write_dorms_excel(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["宿舍号", "床号", "姓名", "籍贯"])
    sheet.append(["501", "01", "张三", "广州"])
    sheet.append(["", "02", "李四", "深圳"])
    workbook.save(path)


def image_bytes(image_format: str = "JPEG", size: tuple[int, int] = (300, 420), color: str = "white") -> bytes:
    buffer = BytesIO()
    Image.new("RGB", size, color).save(buffer, format=image_format)
    return buffer.getvalue()


def test_backend_main_flow(tmp_path: Path, monkeypatch) -> None:
    data_dir = tmp_path / "data"
    monkeypatch.setenv("QXBAN_DATA_DIR", str(data_dir))
    monkeypatch.setenv("QXBAN_DB_PATH", str(data_dir / "qxban.db"))

    from backend.app.config import get_settings
    from backend.app.db import reset_engine_for_tests

    get_settings.cache_clear()
    reset_engine_for_tests()

    import backend.app.main as backend_main

    monkeypatch.setattr(backend_main, "start_photo_background_worker", lambda: None)
    monkeypatch.setattr(backend_main, "stop_photo_background_worker", lambda: None)
    monkeypatch.setattr(backend_main, "photo_background_warmup_state", lambda: {"status": "ready", "error": ""})
    app = backend_main.app

    with TestClient(app) as client:
        health_response = client.get("/api/v1/health")
        assert health_response.status_code == 200
        assert health_response.json()["photo_background_warmup_status"] == "ready"
        template_response = client.get("/api/v1/templates/student-info/download")
        assert template_response.status_code == 200
        assert template_response.content.startswith(b"PK")

        project_response = client.post("/api/v1/projects", json={"name": "测试班级"})
        assert project_response.status_code == 200
        project_id = project_response.json()["id"]

        excel_path = tmp_path / "students.xlsx"
        write_students_excel(excel_path)
        with excel_path.open("rb") as file:
            import_response = client.post(
                f"/api/v1/projects/{project_id}/students/import-excel",
                files={"file": ("students.xlsx", file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert import_response.status_code == 200
        assert import_response.json()["imported_count"] == 2

        students_response = client.get(f"/api/v1/projects/{project_id}/students")
        assert students_response.status_code == 200
        first_student_id = students_response.json()[0]["id"]
        second_student_id = students_response.json()[1]["id"]
        assert {student["class_name"] for student in students_response.json()} == {"测试班级"}

        zhangsan_photo = image_bytes("JPEG", (295, 413), "red")
        lisi_photo = image_bytes("JPEG", (750, 1000), "green")
        stranger_photo = image_bytes("PNG", (320, 420), "blue")
        missing_name_photo = image_bytes("JPEG", (420, 560), "yellow")
        broken_photo = b"not a real image"
        bulk_photo_response = client.post(
            f"/api/v1/projects/{project_id}/photos/import",
            files=[
                ("files", ("01_张三.jpg", zhangsan_photo, "image/jpeg")),
                ("files", ("李四生活照.jpeg", lisi_photo, "image/jpeg")),
                ("files", ("陌生同学.png", stranger_photo, "image/png")),
                ("files", ("06_.jpg", missing_name_photo, "image/jpeg")),
                ("files", ("07_损坏图片.jpg", broken_photo, "image/jpeg")),
            ],
        )
        assert bulk_photo_response.status_code == 200
        bulk_photo_payload = bulk_photo_response.json()
        assert bulk_photo_payload["matched_count"] == 2
        assert bulk_photo_payload["unmatched_count"] == 3
        assert bulk_photo_payload["sequence_numbers"] == [1, 6, 7]
        assert bulk_photo_payload["missing_sequence_numbers"] == [2, 3, 4, 5]
        assert bulk_photo_payload["results"][0]["image_width"] == 295
        assert bulk_photo_payload["results"][0]["image_height"] == 413
        assert bulk_photo_payload["results"][0]["image_format"] == "JPEG"
        assert bulk_photo_payload["results"][4]["status"] == "rejected"
        assert bulk_photo_payload["results"][4]["message"] == "图片文件无法读取，请检查图片是否损坏。"
        assert bulk_photo_payload["results"][4]["image_width"] is None
        assert bulk_photo_payload["results"][4]["file_id"]
        broken_upload_response = client.post(
            f"/api/v1/projects/{project_id}/photos/upload",
            files={"file": ("单张损坏照片.jpg", broken_photo, "image/jpeg")},
        )
        assert broken_upload_response.status_code == 400
        assert broken_upload_response.json()["detail"] == "图片文件无法读取，请检查图片是否损坏。"

        def fake_blue_background(source: Path, target_dir: Path, original_filename: str) -> Path:
            target_dir.mkdir(parents=True, exist_ok=True)
            target = target_dir / f"{Path(original_filename).stem}-blue-test.png"
            Image.new("RGB", (18, 24), "#00acfd").save(target, format="PNG")
            return target

        monkeypatch.setattr(backend_main, "photo_background_available", lambda: True)
        import backend.app.photo_jobs as photo_jobs

        monkeypatch.setattr(photo_jobs, "create_blue_background_photo", fake_blue_background)
        blue_background_response = client.post(
            f"/api/v1/projects/{project_id}/photos/blue-background",
            data={"student_id": str(first_student_id)},
            files={"file": ("单张照片.jpg", zhangsan_photo, "image/jpeg")},
        )
        assert blue_background_response.status_code == 200
        blue_background_job = blue_background_response.json()
        assert blue_background_job["status"] == "queued"
        assert blue_background_job["student_id"] == first_student_id
        assert blue_background_job["queue_ahead_count"] == 0
        second_blue_background_response = client.post(
            f"/api/v1/projects/{project_id}/photos/blue-background",
            data={"student_id": str(second_student_id)},
            files={"file": ("第二张照片.jpg", lisi_photo, "image/jpeg")},
        )
        assert second_blue_background_response.status_code == 200
        second_blue_background_job = second_blue_background_response.json()
        assert second_blue_background_job["status"] == "queued"
        assert second_blue_background_job["queue_ahead_count"] == 1
        from backend.app.photo_jobs import _claim_next_photo_background_job, process_photo_background_job

        assert _claim_next_photo_background_job() == blue_background_job["id"]
        process_photo_background_job(blue_background_job["id"])
        assert _claim_next_photo_background_job() == second_blue_background_job["id"]
        process_photo_background_job(second_blue_background_job["id"])
        completed_job_response = client.get(
            f"/api/v1/projects/{project_id}/photos/blue-background/jobs/{blue_background_job['id']}"
        )
        assert completed_job_response.status_code == 200
        completed_job = completed_job_response.json()
        assert completed_job["status"] == "succeeded"
        assert completed_job["result"]["filename"].endswith("-blue-test.png")
        blue_background_download = client.get(completed_job["result"]["download_url"])
        assert blue_background_download.status_code == 200
        assert blue_background_download.headers["content-type"].startswith("image/png")
        with Image.open(BytesIO(blue_background_download.content)) as image:
            assert image.size == (18, 24)
            assert image.getpixel((0, 0)) == (0, 172, 253)
        activities_before_expiry = client.get(f"/api/v1/projects/{project_id}/activities?limit=50")
        assert activities_before_expiry.status_code == 200
        blue_success_activities = [
            activity
            for activity in activities_before_expiry.json()
            if activity["message"] == "张三蓝底设置成功"
        ]
        assert blue_success_activities
        assert blue_success_activities[0]["detail"]["highlight"] is True
        from sqlmodel import Session

        from backend.app.db import get_engine
        from backend.app.files import absolute_path
        from backend.app.models import FileRecord, PhotoBackgroundJobRecord

        with Session(get_engine()) as session:
            jobs_before_expiry = [
                session.get(PhotoBackgroundJobRecord, job_id)
                for job_id in (blue_background_job["id"], second_blue_background_job["id"])
            ]
            assert all(job is not None and job.completed_at is not None for job in jobs_before_expiry)
            job_file_paths = [
                absolute_path(file_record.path)
                for job in jobs_before_expiry
                for file_id in (job.input_file_id, job.result_file_id)
                if file_id
                if (file_record := session.get(FileRecord, file_id)) is not None
            ]
            expires_at = max(job.completed_at for job in jobs_before_expiry if job.completed_at is not None) + photo_jobs.PHOTO_BACKGROUND_RESULT_TTL + timedelta(seconds=1)
        assert photo_jobs.cleanup_expired_photo_background_jobs(expires_at) == 2
        assert client.get(f"/api/v1/projects/{project_id}/photos/blue-background/jobs/{blue_background_job['id']}").status_code == 404
        assert client.get(completed_job["result"]["download_url"]).status_code == 404
        assert all(not path.exists() for path in job_file_paths)
        activities_after_expiry = client.get(f"/api/v1/projects/{project_id}/activities?limit=50")
        assert activities_after_expiry.status_code == 200
        assert not any(activity["kind"] == "photo_blue_background" for activity in activities_after_expiry.json())
        assert len(bulk_photo_payload["sequence_gap_students"]) == 1
        assert bulk_photo_payload["sequence_gap_students"][0]["index"] == 2
        assert bulk_photo_payload["sequence_gap_students"][0]["name"] == "李四"
        assert bulk_photo_payload["sequence_gap_students"][0]["photo_bound"] is True
        assert bulk_photo_payload["report_file_id"]
        assert bulk_photo_payload["report_download_url"].startswith("/api/v1/files/")
        report_response = client.get(bulk_photo_payload["report_download_url"])
        assert report_response.status_code == 200
        report_workbook = load_workbook(BytesIO(report_response.content), data_only=True)
        report_sheet = report_workbook["照片导入核对"]
        assert report_sheet["A1"].value == "照片导入核对清单"
        assert report_sheet["B2"].value == 2
        assert report_sheet["B3"].value == 3
        assert report_sheet["B5"].value == "02、03、04、05"
        assert report_sheet["B8"].value == "01_张三.jpg"
        assert report_sheet["H8"].value == "295 x 413"
        assert report_sheet["I8"].value == "JPEG"
        assert "学生照片覆盖" in report_workbook.sheetnames
        assert "缺照片学生" in report_workbook.sheetnames
        assert "文件序号缺口" in report_workbook.sheetnames
        assert "图片规格汇总" in report_workbook.sheetnames
        image_summary_sheet = report_workbook["图片规格汇总"]
        assert [image_summary_sheet.cell(1, column).value for column in range(1, 5)] == ["图片尺寸", "图片格式", "数量", "示例文件"]
        image_summary = {
            (image_summary_sheet.cell(row, 1).value, image_summary_sheet.cell(row, 2).value): image_summary_sheet.cell(row, 3).value
            for row in range(2, image_summary_sheet.max_row + 1)
        }
        assert image_summary[("295 x 413", "JPEG")] == 1
        assert image_summary[("750 x 1000", "JPEG")] == 1
        assert image_summary[("320 x 420", "PNG")] == 1
        assert image_summary[("420 x 560", "JPEG")] == 1
        assert image_summary[("无法读取", "未知")] == 1
        sequence_gap_sheet = report_workbook["文件序号缺口"]
        assert [sequence_gap_sheet.cell(2, column).value for column in range(1, 8)] == [
            "02",
            "李四",
            bulk_photo_payload["sequence_gap_students"][0]["student_id"],
            "深圳",
            "测试班级",
            "已绑定",
            "模糊",
        ]
        coverage_sheet = report_workbook["学生照片覆盖"]
        assert [coverage_sheet.cell(row, 2).value for row in (2, 3)] == ["张三", "李四"]
        assert [coverage_sheet.cell(row, 5).value for row in (2, 3)] == ["已绑定", "已绑定"]
        missing_sheet = report_workbook["缺照片学生"]
        assert missing_sheet.max_row == 1
        assert [item["status"] for item in bulk_photo_payload["results"]] == ["strict", "fuzzy", "unmatched", "unmatched", "rejected"]
        assert bulk_photo_payload["results"][3]["message"] == "文件名缺少学生姓名"
        assert bulk_photo_payload["results"][3]["sequence_number"] == 6
        assert bulk_photo_payload["results"][3]["suggested_student_id"] is None
        assert bulk_photo_payload["results"][3]["file_id"]
        assert bulk_photo_payload["results"][3]["download_url"].startswith("/api/v1/files/")
        unmatched_photo_download = client.get(bulk_photo_payload["results"][3]["download_url"])
        assert unmatched_photo_download.status_code == 200
        assert unmatched_photo_download.content == missing_name_photo
        students_after_bulk_photo = client.get(f"/api/v1/projects/{project_id}/students").json()
        photo_status_by_name = {student["name"]: student for student in students_after_bulk_photo}
        assert photo_status_by_name["张三"]["photo_match_status"] == "strict"
        assert photo_status_by_name["张三"]["photo_url"]
        assert photo_status_by_name["李四"]["photo_match_status"] == "fuzzy"
        assert photo_status_by_name["李四"]["photo_url"]
        zhangsan_bound_photo = client.get(photo_status_by_name["张三"]["photo_url"])
        assert zhangsan_bound_photo.status_code == 200
        assert zhangsan_bound_photo.headers["content-type"].startswith("image/png")
        with Image.open(BytesIO(zhangsan_bound_photo.content)) as image:
            assert image.size == (160, 210)
            assert image.format == "PNG"
        photos_zip_response = client.post(f"/api/v1/projects/{project_id}/photos/export-zip")
        assert photos_zip_response.status_code == 200
        photos_zip_download = client.get(photos_zip_response.json()["download_url"])
        assert photos_zip_download.status_code == 200
        with ZipFile(BytesIO(photos_zip_download.content)) as archive:
            assert sorted(archive.namelist()) == ["01_张三.png", "02_李四.png"]
            with archive.open("01_张三.png") as zipped_photo:
                with Image.open(BytesIO(zipped_photo.read())) as image:
                    assert image.size == (160, 210)
                    assert image.format == "PNG"
        attach_unmatched_response = client.post(
            f"/api/v1/projects/{project_id}/students/{first_student_id}/photo",
            json={"file_id": bulk_photo_payload["results"][3]["file_id"]},
        )
        assert attach_unmatched_response.status_code == 200
        assert attach_unmatched_response.json()["photo_match_status"] == "manual"
        assert attach_unmatched_response.json()["photo_url"]
        manual_bound_photo = client.get(attach_unmatched_response.json()["photo_url"])
        assert manual_bound_photo.status_code == 200
        with Image.open(BytesIO(manual_bound_photo.content)) as image:
            assert image.size == (160, 210)

        status_report_response = client.post(f"/api/v1/projects/{project_id}/photos/status-report")
        assert status_report_response.status_code == 200
        status_report_download = client.get(status_report_response.json()["download_url"])
        assert status_report_download.status_code == 200
        status_report_workbook = load_workbook(BytesIO(status_report_download.content), data_only=True)
        status_sheet = status_report_workbook["照片状态清单"]
        assert status_sheet["A1"].value == "学生照片状态清单"
        assert status_sheet["B2"].value == 2
        assert status_sheet["B3"].value == 2
        assert status_sheet["B4"].value == 0
        assert [status_sheet.cell(row, 2).value for row in (7, 8)] == ["张三", "李四"]
        assert [status_sheet.cell(row, 5).value for row in (7, 8)] == ["已绑定", "已绑定"]
        assert [status_sheet.cell(row, 6).value for row in (7, 8)] == ["手动", "模糊"]
        assert status_report_workbook["缺照片学生"].max_row == 1

        clear_photo_response = client.delete(f"/api/v1/projects/{project_id}/students/{first_student_id}/photo")
        assert clear_photo_response.status_code == 200
        assert clear_photo_response.json()["photo_match_status"] == "missing"
        assert clear_photo_response.json()["photo_url"] is None
        cleared_status_report_response = client.post(f"/api/v1/projects/{project_id}/photos/status-report")
        assert cleared_status_report_response.status_code == 200
        cleared_status_report_download = client.get(cleared_status_report_response.json()["download_url"])
        assert cleared_status_report_download.status_code == 200
        cleared_status_workbook = load_workbook(BytesIO(cleared_status_report_download.content), data_only=True)
        cleared_status_sheet = cleared_status_workbook["照片状态清单"]
        assert cleared_status_sheet["B3"].value == 1
        assert cleared_status_sheet["B4"].value == 1
        cleared_missing_sheet = cleared_status_workbook["缺照片学生"]
        assert cleared_missing_sheet.max_row == 2
        assert cleared_missing_sheet["B2"].value == "张三"

        dorm_path = tmp_path / "dorms.xlsx"
        write_dorms_excel(dorm_path)
        with dorm_path.open("rb") as file:
            dorm_import_response = client.post(
                f"/api/v1/projects/{project_id}/dorms/import-excel",
                files={"file": ("dorms.xlsx", file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert dorm_import_response.status_code == 200
        assert dorm_import_response.json()["room_count"] == 1
        assert dorm_import_response.json()["matched_count"] == 2
        recent_dorms_response = client.get(f"/api/v1/projects/{project_id}/dorms/recent")
        assert recent_dorms_response.status_code == 200
        recent_dorms = recent_dorms_response.json()
        assert recent_dorms[0]["filename"] == "dorms.xlsx"
        restore_recent_response = client.post(
            f"/api/v1/projects/{project_id}/dorms/recent/{recent_dorms[0]['file_id']}/restore"
        )
        assert restore_recent_response.status_code == 200
        assert restore_recent_response.json()["matched_count"] == 2
        dorm_export_response = client.post(
            f"/api/v1/projects/{project_id}/dorms/export-excel",
            json={"title": "测试宿舍安排表"},
        )
        assert dorm_export_response.status_code == 200
        recent_after_export = client.get(f"/api/v1/projects/{project_id}/dorms/recent").json()
        assert recent_after_export[0]["filename"].startswith("宿舍安排表-")
        restore_export_response = client.post(
            f"/api/v1/projects/{project_id}/dorms/recent/{recent_after_export[0]['file_id']}/restore"
        )
        assert restore_export_response.status_code == 200
        assert restore_export_response.json()["matched_count"] == 2

        patch_response = client.patch(
            f"/api/v1/projects/{project_id}/students/{first_student_id}",
            json={"hometown": "广州花都"},
        )
        assert patch_response.status_code == 200
        assert patch_response.json()["hometown"] == "广州花都"

        create_response = client.post(
            f"/api/v1/projects/{project_id}/students",
            json={"name": "王五", "hometown": "佛山", "class_name": "一班"},
        )
        assert create_response.status_code == 200
        assert create_response.json()["index"] == 3
        assert create_response.json()["name"] == "王五"
        assert create_response.json()["class_name"] == "测试班级"

        photo_response = client.post(
            f"/api/v1/projects/{project_id}/photos/upload",
            files={"file": ("wangwu.png", image_bytes("PNG", (320, 420), "purple"), "image/png")},
        )
        assert photo_response.status_code == 200
        attach_response = client.post(
            f"/api/v1/projects/{project_id}/students/{create_response.json()['id']}/photo",
            json={"file_id": photo_response.json()["file_id"]},
        )
        assert attach_response.status_code == 200
        assert attach_response.json()["photo_match_status"] == "manual"
        assert attach_response.json()["photo_url"]

        seats_response = client.get(f"/api/v1/projects/{project_id}/seats")
        assert seats_response.status_code == 200
        assert seats_response.json()["layout"]["seats"]
        assert "王五" in seats_response.json()["layout"]["student_names"]

        delete_response = client.delete(f"/api/v1/projects/{project_id}/students/{first_student_id}")
        assert delete_response.status_code == 200
        students_after_delete = client.get(f"/api/v1/projects/{project_id}/students").json()
        assert [student["index"] for student in students_after_delete] == [1, 2]
        assert [student["name"] for student in students_after_delete] == ["李四", "王五"]
        seats_after_delete = client.get(f"/api/v1/projects/{project_id}/seats").json()
        assert seats_after_delete["layout"]["student_names"] == ["李四", "王五"]

        second_project_response = client.post("/api/v1/projects", json={"name": "测试二班"})
        assert second_project_response.status_code == 200
        second_project_id = second_project_response.json()["id"]
        second_create_response = client.post(
            f"/api/v1/projects/{second_project_id}/students",
            json={"name": "赵六", "hometown": "珠海", "class_name": "二班"},
        )
        assert second_create_response.status_code == 200
        second_students = client.get(f"/api/v1/projects/{second_project_id}/students").json()
        assert [student["name"] for student in second_students] == ["赵六"]
        first_students = client.get(f"/api/v1/projects/{project_id}/students").json()
        assert [student["name"] for student in first_students] == ["李四", "王五"]
        second_seats = client.get(f"/api/v1/projects/{second_project_id}/seats").json()
        assert second_seats["layout"]["student_names"] == ["赵六"]
        first_seats = client.get(f"/api/v1/projects/{project_id}/seats").json()
        assert first_seats["layout"]["student_names"] == ["李四", "王五"]
        saved_layout = first_seats["layout"]
        saved_layout["title_text"] = "一班保存后的座位表"
        saved_layout["left"] = 188
        saved_layout["top"] = 166
        saved_layout["scale_percent"] = 60
        saved_layout["seats"][0]["x"] = 188
        saved_layout["seats"][0]["y"] = 166
        save_seats_response = client.put(f"/api/v1/projects/{project_id}/seats", json={"layout": saved_layout})
        assert save_seats_response.status_code == 200
        reloaded_seats = client.get(f"/api/v1/projects/{project_id}/seats").json()["layout"]
        assert reloaded_seats["title_text"] == "一班保存后的座位表"
        assert reloaded_seats["left"] == 188
        assert reloaded_seats["scale_percent"] == 60
        assert reloaded_seats["seats"][0]["x"] == 188
        second_seats_after_save = client.get(f"/api/v1/projects/{second_project_id}/seats").json()["layout"]
        assert second_seats_after_save["title_text"] == "班级座位表"

        export_response = client.post(f"/api/v1/projects/{project_id}/seats/export-pdf")
        assert export_response.status_code == 200
        download_response = client.get(export_response.json()["download_url"])
        assert download_response.status_code == 200
        assert download_response.content.startswith(b"%PDF")

        rename_response = client.patch(f"/api/v1/projects/{project_id}", json={"name": "测试班级已改名"})
        assert rename_response.status_code == 200
        assert rename_response.json()["name"] == "测试班级已改名"
        renamed_students = client.get(f"/api/v1/projects/{project_id}/students").json()
        assert {student["class_name"] for student in renamed_students} == {"测试班级已改名"}
        activities_response = client.get(f"/api/v1/projects/{project_id}/activities?limit=50")
        assert activities_response.status_code == 200
        activity_messages = [activity["message"] for activity in activities_response.json()]
        assert any("创建班级" in message for message in activity_messages)
        assert any("导入学生名单" in message for message in activity_messages)
        assert any("导出座位表" in message for message in activity_messages)
        assert any("班级改名" in message for message in activity_messages)
        blue_success_activities = [
            activity
            for activity in activities_response.json()
            if activity["message"] == "张三蓝底设置成功"
        ]
        assert not blue_success_activities


def test_photo_import_suggests_student_by_sequence(tmp_path: Path, monkeypatch) -> None:
    data_dir = tmp_path / "data"
    monkeypatch.setenv("QXBAN_DATA_DIR", str(data_dir))
    monkeypatch.setenv("QXBAN_DB_PATH", str(data_dir / "qxban.db"))
    monkeypatch.setenv("QXBAN_PHOTO_BACKGROUND_PREWARM_ENABLED", "0")

    from backend.app.config import get_settings
    from backend.app.db import reset_engine_for_tests

    get_settings.cache_clear()
    reset_engine_for_tests()

    from backend.app.main import app

    with TestClient(app) as client:
        project_response = client.post("/api/v1/projects", json={"name": "序号建议班"})
        assert project_response.status_code == 200
        project_id = project_response.json()["id"]

        excel_path = tmp_path / "students.xlsx"
        write_numbered_students_excel(excel_path, ["一号", "二号", "三号", "四号", "五号", "六号"])
        with excel_path.open("rb") as file:
            import_response = client.post(
                f"/api/v1/projects/{project_id}/students/import-excel",
                files={"file": ("students.xlsx", file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert import_response.status_code == 200

        photo_response = client.post(
            f"/api/v1/projects/{project_id}/photos/import",
            files=[
                ("files", ("06_.jpg", image_bytes("JPEG", (420, 560), "yellow"), "image/jpeg")),
            ],
        )
        assert photo_response.status_code == 200
        payload = photo_response.json()
        assert payload["matched_count"] == 0
        result = payload["results"][0]
        assert result["status"] == "unmatched"
        assert result["sequence_number"] == 6
        assert result["suggested_student_name"] == "六号"
        assert result["suggested_student_index"] == 6
        assert result["suggested_student_id"]

        duplicate_named_response = client.post(
            f"/api/v1/projects/{project_id}/photos/import",
            files=[
                ("files", ("06_.jpg", image_bytes("JPEG", (420, 560), "yellow"), "image/jpeg")),
                ("files", ("06_六号.jpg", image_bytes("JPEG", (420, 560), "green"), "image/jpeg")),
            ],
        )
        assert duplicate_named_response.status_code == 200
        duplicate_named_payload = duplicate_named_response.json()
        assert duplicate_named_payload["results"][0]["status"] == "unmatched"
        assert duplicate_named_payload["results"][0]["sequence_number"] == 6
        assert duplicate_named_payload["results"][0]["suggested_student_id"] is None
