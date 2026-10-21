from io import BytesIO
import sqlite3
from contextlib import contextmanager
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlmodel import Session, select


def excel_bytes(rows):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    data = BytesIO()
    workbook.save(data)
    return data.getvalue()


@contextmanager
def isolated_app(tmp_path, monkeypatch):
    monkeypatch.setenv('QXBAN_DATA_DIR', str(tmp_path))
    monkeypatch.setenv('QXBAN_DB_PATH', str(tmp_path / 'qxban.db'))
    from backend.app.config import get_settings
    from backend.app.db import reset_engine_for_tests
    import backend.app.main as main
    get_settings.cache_clear()
    reset_engine_for_tests()
    monkeypatch.setattr(main, 'start_photo_background_worker', lambda: None)
    monkeypatch.setattr(main, 'stop_photo_background_worker', lambda: None)
    monkeypatch.setattr(main, 'start_registration_worker', lambda: None)
    monkeypatch.setattr(main, 'stop_registration_worker', lambda: None)
    try:
        with TestClient(main.app) as client:
            yield client, main.app
    finally:
        reset_engine_for_tests()
        get_settings.cache_clear()


def setup(client):
    response = client.post('/api/v1/auth/setup', json={'username': 'admin', 'password': 'AdminPass2026!'})
    assert response.status_code == 200, response.text
    return response


def test_accounts_permissions_reset_and_import(tmp_path, monkeypatch):
    with isolated_app(tmp_path, monkeypatch) as (admin, app):
        assert admin.get('/api/v1/projects').status_code == 401
        assert admin.get('/api/v1/admin/users').status_code == 401
        assert admin.get('/api/v1/auth/status').json() == {'needs_setup': True}
        response = setup(admin)
        cookie = response.headers['set-cookie']
        assert 'HttpOnly' in cookie and 'SameSite=strict' in cookie
        assert 'password' not in response.json()
        assert admin.post('/api/v1/auth/setup', json={'username': 'other', 'password': 'OtherPass2026!'}).status_code == 409
        assert admin.post('/api/v1/admin/users', json={'username': 'teacher1', 'password': 'TeacherPass2026!', 'role': 'admin'}).json()['role'] == 'user'
        users = admin.get('/api/v1/admin/users').json()
        teacher_id = next(row['id'] for row in users if row['username'] == 'teacher1')
        assert admin.post('/api/v1/admin/users', json={'username': 'teacher1', 'password': 'TeacherPass2026!'}).status_code == 409
        assert admin.post('/api/v1/admin/users', json={'username': 'teacher2', 'password': 'short'}).status_code == 400
        root_project = admin.post('/api/v1/projects', json={'name': '管理员班级'}).json()['id']
        source = excel_bytes([['账号', '姓名', '初始密码'], ['teacher2', '陈老师', 'TeacherPass2026!'], ['teacher3', '张老师', 'TeacherPass2026!']])
        imported = admin.post('/api/v1/admin/users/import', files={'file': ('users.xlsx', source)})
        assert imported.status_code == 200, imported.text
        assert imported.json()['imported_count'] == 2
        invalid = excel_bytes([['账号', '姓名', '初始密码'], ['teacher4', '李老师', 'TeacherPass2026!'], ['teacher2', '陈老师', 'TeacherPass2026!']])
        assert admin.post('/api/v1/admin/users/import', files={'file': ('users.xlsx', invalid)}).status_code == 400
        assert 'teacher4' not in [row['username'] for row in admin.get('/api/v1/admin/users').json()]
        assert admin.get('/api/v1/admin/users/template').content.startswith(b'PK')
        # A client without the admin's cookie logs in as an ordinary user.
        with TestClient(app) as teacher:
            login = teacher.post('/api/v1/auth/login', json={'username': 'teacher1', 'password': 'TeacherPass2026!'})
            assert login.status_code == 200
            assert teacher.get('/api/v1/projects').status_code == 403
            assert teacher.post('/api/v1/auth/change-password', json={'current_password': 'TeacherPass2026!', 'new_password': 'MyTeacherPass2026!'}).status_code == 200
            assert teacher.get('/api/v1/admin/users').status_code == 403
            assert teacher.post('/api/v1/admin/users', json={'username': 'intruder', 'password': 'SomePassword2026!'}).status_code == 403
            assert teacher.get(f'/api/v1/projects/{root_project}').status_code == 404
            assert teacher.get(f'/api/v1/projects/{root_project}/students').status_code == 404
            project = teacher.post('/api/v1/projects', json={'name': '教师班级'}).json()['id']
            assert [item['id'] for item in teacher.get('/api/v1/projects').json()] == [project]
            assert {item['id'] for item in admin.get('/api/v1/projects').json()} == {root_project, project}
            export = admin.post(f'/api/v1/projects/{root_project}/students/export-excel').json()
            assert teacher.get(export['download_url']).status_code == 404
            assert teacher.post('/api/v1/projects', json={'name': '跨站'}, headers={'origin': 'https://attacker.example'}).status_code == 403
            assert admin.post(f'/api/v1/admin/users/{teacher_id}/reset-password', json={'password': 'ResetTeacher2026!'}).status_code == 200
            assert teacher.get('/api/v1/auth/me').status_code == 401
            assert teacher.post('/api/v1/auth/login', json={'username': 'teacher1', 'password': 'MyTeacherPass2026!'}).status_code == 401
            assert teacher.post('/api/v1/auth/login', json={'username': 'teacher1', 'password': 'ResetTeacher2026!'}).json()['must_change_password'] is True
            assert teacher.post('/api/v1/auth/change-password', json={'current_password': 'ResetTeacher2026!', 'new_password': 'FinalTeacher2026!'}).status_code == 200
            assert admin.patch(f'/api/v1/admin/users/{teacher_id}', json={'is_active': False}).status_code == 200
            assert teacher.get('/api/v1/projects').status_code == 401
            assert teacher.post('/api/v1/auth/login', json={'username': 'teacher1', 'password': 'FinalTeacher2026!'}).status_code == 401
            assert admin.patch(f'/api/v1/admin/users/{teacher_id}', json={'is_active': True}).status_code == 200
        from backend.app.db import get_engine
        from backend.app.models import UserRecord, LoginSessionRecord
        with Session(get_engine()) as session:
            for user in session.exec(select(UserRecord)):
                assert user.password_hash.startswith('scrypt$')
                assert user.password_hash != 'AdminPass2026!'
            assert all(len(row.id) == 64 for row in session.exec(select(LoginSessionRecord)))
        assert admin.post('/api/v1/auth/logout').status_code == 200
        assert admin.get('/api/v1/projects').status_code == 401


def test_source_area_roundtrip_and_old_import(tmp_path, monkeypatch):
    with isolated_app(tmp_path, monkeypatch) as (client, _app):
        setup(client)
        project_id = client.post('/api/v1/projects', json={'name': '测试班级'}).json()['id']
        path = f'/api/v1/projects/{project_id}'
        source = excel_bytes([['序号', '姓名', '籍贯', '班级', '送生区'], [1, '张三', '广州', '测试班级', '白云区']])
        response = client.post(path + '/students/import-excel', files={'file': ('students.xlsx', source)})
        assert response.status_code == 200, response.text
        student = response.json()['students'][0]
        assert student['source_area'] == '白云区'
        changed = client.patch(path + f"/students/{student['id']}", json={'source_area': '越秀区'})
        assert changed.json()['source_area'] == '越秀区'
        new = client.post(path + '/students', json={'name': '李四', 'source_area': '海珠区'})
        assert new.json()['source_area'] == '海珠区'
        export = client.post(path + '/students/export-excel').json()
        workbook = load_workbook(BytesIO(client.get(export['download_url']).content))
        assert [cell.value for cell in workbook.active[1]] == ['序号', '姓名', '籍贯', '班级', '送生区']
        assert workbook.active['E2'].value == '越秀区'
        client.put(path + '/dorms', json={'rooms': [{'number': '101', 'beds': [str(student['id']), None]}]})
        export = client.post(path + '/dorms/export-excel', json={'title': '测试宿舍'}).json()
        workbook = load_workbook(BytesIO(client.get(export['download_url']).content))
        assert workbook.active['E2'].value == '送生区'
        assert workbook.active['E3'].value == '越秀区'
        export = client.post(path + '/stickers/export-pdf', json={'sticker_kind': 'door'}).json()
        assert client.get(export['download_url']).content.startswith(b'%PDF')
        empty = excel_bytes([['序号', '姓名', '籍贯', '班级', '送生区']])
        assert client.post(path + '/students/import-excel', files={'file': ('empty.xlsx', empty)}).status_code == 400
        assert len(client.get(path + '/students').json()) == 2
        old = excel_bytes([['序号', '姓名', '籍贯', '班级'], [1, '王五', '广州', '测试班级']])
        response = client.post(path + '/students/import-excel', files={'file': ('old.xlsx', old)})
        assert response.status_code == 200, response.text
        assert response.json()['students'][0]['source_area'] == ''
        template = load_workbook(BytesIO(client.get('/api/v1/templates/student-info/download').content))
        assert template.active['E1'].value == '送生区'


def test_additive_database_migration(tmp_path, monkeypatch):
    db = sqlite3.connect(tmp_path / 'qxban.db')
    db.executescript("""
      CREATE TABLE projectrecord (id INTEGER PRIMARY KEY, name TEXT NOT NULL, created_at DATETIME NOT NULL);
      INSERT INTO projectrecord VALUES (1, '历史班级', '2026-01-01 00:00:00');
      CREATE TABLE studentrecord (id INTEGER PRIMARY KEY, project_id INTEGER NOT NULL, student_index INTEGER NOT NULL,
        name TEXT NOT NULL, hometown TEXT NOT NULL, class_name TEXT NOT NULL, photo_path TEXT NOT NULL, photo_match_status TEXT NOT NULL);
      INSERT INTO studentrecord VALUES (1, 1, 1, '历史学生', '广州', '历史班级', '', 'unmatched');
    """)
    db.close()
    with isolated_app(tmp_path, monkeypatch) as (client, _app):
        setup(client)
        students = client.get('/api/v1/projects/1/students').json()
        assert students[0]['name'] == '历史学生'
        assert students[0]['source_area'] == ''
        with sqlite3.connect(tmp_path / 'qxban.db') as db:
            assert db.execute('select owner_id from projectrecord where id=1').fetchone()[0] is not None
            assert db.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
